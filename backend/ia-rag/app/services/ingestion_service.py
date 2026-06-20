import json
from pathlib import Path

from docx import Document
from pypdf import PdfReader

MAX_FILE_SIZE_BYTES = 20 * 1024 * 1024  # 20 MB


class IngestionService:
    def __init__(self) -> None:
        self.base_path = Path(__file__).resolve().parents[2]
        self.books_path = self.base_path / "data" / "books"
        self.processed_path = self.base_path / "data" / "processed"

    def validate_file_size(self, content: bytes) -> None:
        if len(content) > MAX_FILE_SIZE_BYTES:
            raise ValueError(
                f"El archivo supera el límite de {MAX_FILE_SIZE_BYTES // (1024 * 1024)} MB"
            )

    def extract_pdf_text(self, pdf_filename: str) -> str:
        pdf_path = self.books_path / pdf_filename
        if not pdf_path.exists():
            raise FileNotFoundError(f"No se encontró el archivo: {pdf_path}")
        reader = PdfReader(str(pdf_path))
        pages = [page.extract_text() or "" for page in reader.pages]
        return "\n\n".join(p.strip() for p in pages if p.strip())

    def extract_docx_text(self, docx_filename: str) -> str:
        docx_path = self.books_path / docx_filename
        if not docx_path.exists():
            raise FileNotFoundError(f"No se encontró el archivo: {docx_path}")
        doc = Document(str(docx_path))
        parts = []
        for p in doc.paragraphs:
            if p.text.strip():
                parts.append(p.text.strip())
        for table in doc.tables:
            for row in table.rows:
                # dict.fromkeys preserves order and deduplicates merged cells
                cells = list(dict.fromkeys(
                    c.text.strip() for c in row.cells if c.text.strip()
                ))
                if cells:
                    parts.append(" | ".join(cells))
        return "\n\n".join(parts)

    def save_text_output(self, output_filename: str, text: str) -> Path:
        self.processed_path.mkdir(parents=True, exist_ok=True)
        output_path = self.processed_path / output_filename
        output_path.write_text(text, encoding="utf-8")
        return output_path

    def ingest_pdf(self, pdf_filename: str) -> Path:
        text = self.extract_pdf_text(pdf_filename)
        return self.save_text_output(Path(pdf_filename).stem + ".txt", text)

    def ingest_docx(self, docx_filename: str) -> Path:
        text = self.extract_docx_text(docx_filename)
        return self.save_text_output(Path(docx_filename).stem + ".txt", text)

    def ingest_file(self, filename: str) -> Path:
        suffix = Path(filename).suffix.lower()
        if suffix == ".pdf":
            return self.ingest_pdf(filename)
        if suffix == ".docx":
            return self.ingest_docx(filename)
        raise ValueError(f"Formato no soportado: {suffix}. Use .pdf o .docx")

    def update_metadata(self, filename: str, label: str) -> None:
        metadata_path = self.base_path / "data" / "books_metadata.json"
        books = []
        if metadata_path.exists():
            books = json.loads(metadata_path.read_text(encoding="utf-8"))
        txt_filename = Path(filename).stem + ".txt"
        if any(b["filename"] == txt_filename for b in books):
            return
        books.append({"filename": txt_filename, "label": label, "version": "new"})
        metadata_path.write_text(
            json.dumps(books, ensure_ascii=False, indent=2), encoding="utf-8"
        )

    def remove_book(self, txt_filename: str) -> None:
        metadata_path = self.base_path / "data" / "books_metadata.json"
        books = []
        if metadata_path.exists():
            books = json.loads(metadata_path.read_text(encoding="utf-8"))
        books = [b for b in books if b["filename"] != txt_filename]
        metadata_path.write_text(
            json.dumps(books, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        txt_path = self.processed_path / txt_filename
        if txt_path.exists():
            txt_path.unlink()
        stem = Path(txt_filename).stem
        for original in self.books_path.iterdir():
            if original.stem == stem and original.suffix.lower() in (".pdf", ".docx"):
                original.unlink()
                break


ingestion_service = IngestionService()
